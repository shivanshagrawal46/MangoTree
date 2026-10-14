"""Exports for an ARV analysis: a PDF anyone can forward (the valuation, the
comps grid with every adjustment, the market read, the deal math and the
analyst's review) and the MAO Rehab Calculator as a spreadsheet laid out like
the team's own CSV, with live formulas so the numbers can be worked by hand.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional


def _m(v: Any, dash: str = "—") -> str:
    try:
        if v is None:
            return dash
        return "${:,.0f}".format(float(v))
    except (TypeError, ValueError):
        return dash


def _n(v: Any, dash: str = "—", nd: int = 0) -> str:
    try:
        if v is None:
            return dash
        return ("{:,.%df}" % nd).format(float(v))
    except (TypeError, ValueError):
        return dash


def _esc(s: Any) -> str:
    return str(s if s is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _md(s: str) -> str:
    """Just enough markdown for reportlab paragraphs: bold, bullets, breaks."""
    t = _esc(s)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?m)^\s*[-*]\s+", "• ", t)
    t = re.sub(r"(?m)^#+\s*", "", t)
    return t.replace("\n", "<br/>")


# ================================================================== PDF
def pdf(doc: Dict[str, Any], *, user_name: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    pdoc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=16 * mm,
                             title="ARV analysis", author="MangoTree")
    ss = getSampleStyleSheet()
    st = {"h1": ParagraphStyle("h1", parent=ss["Title"], fontSize=17, leading=21, spaceAfter=2, alignment=0),
          "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, leading=15, spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#1f6f5f")),
          "p": ParagraphStyle("p", parent=ss["BodyText"], fontSize=9.2, leading=12.5),
          "small": ParagraphStyle("small", parent=ss["BodyText"], fontSize=7.6, leading=10, textColor=colors.HexColor("#6b6b76")),
          "big": ParagraphStyle("big", parent=ss["BodyText"], fontSize=22, leading=26, textColor=colors.HexColor("#1f6f5f"))}
    P = Paragraph
    story: List[Any] = []
    geo = doc.get("geocode") or {}
    subj = doc.get("subject") or {}
    v = doc.get("valuation") or {}
    rec = v.get("reconciliation") or {}
    conf = v.get("confidence") or {}
    an = doc.get("analysis") or {}
    uw = doc.get("underwriting") or {}
    rb = doc.get("rehab_budget") or {}

    story.append(P("After Repair Value — " + _esc(geo.get("formatted") or subj.get("address")), st["h1"]))
    story.append(P("MangoTree · RKB Consulting Group · %s%s · analysis %s" % (datetime.now(timezone.utc).strftime("%B %d, %Y"),
                   (" · prepared for " + _esc(user_name)) if user_name else "", _esc(doc.get("analysis_id"))), st["small"]))
    story.append(Spacer(1, 6))
    facts = " · ".join(x for x in [
        (subj.get("property_type") or "").replace("_", " ") or None,
        ("%s bd" % subj["beds"]) if subj.get("beds") is not None else None, ("%s ba" % subj["baths"]) if subj.get("baths") is not None else None,
        ("%s sqft" % _n(subj.get("sqft"))) if subj.get("sqft") else None, ("lot %s sqft" % _n(subj.get("lot_sqft"))) if subj.get("lot_sqft") else None,
        ("built %s" % subj["year_built"]) if subj.get("year_built") else None,
        ("as repaired %s" % subj.get("condition_after")) if subj.get("condition_after") else None] if x)
    story.append(P(facts, st["p"]))
    story.append(Spacer(1, 8))
    head = [[P("<b>ARV (sales comparison)</b>", st["p"]), P("<b>Range</b>", st["p"]), P("<b>Confidence</b>", st["p"]), P("<b>Analyst opinion</b>", st["p"]), P("<b>Recommended MAO</b>", st["p"])],
            [P(_m(v.get("arv")), st["big"]), P("%s – %s" % (_m(v.get("arv_low")), _m(v.get("arv_high"))), st["p"]),
             P("%s (%s/100, ±%s%%)" % (_esc(conf.get("label")), conf.get("score"), conf.get("band_pct")), st["p"]),
             P("%s<br/>%s – %s" % (_m(an.get("arv_opinion")), _m(an.get("arv_low")), _m(an.get("arv_high"))) if an else "—", st["p"]),
             P(_m((uw.get("mao") or {}).get("recommended")), st["p"])]]
    t = Table(head, colWidths=[42 * mm, 36 * mm, 38 * mm, 36 * mm, 30 * mm])
    t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d3cfc5")), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e6e3dc")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efede8")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(t)
    if an.get("headline"):
        story.append(Spacer(1, 6))
        story.append(P("<b>%s</b>" % _esc(an["headline"]), st["p"]))

    # ------------------------------------------------------------- comps
    story.append(P("Comparable sales", st["h2"]))
    comps = [c for c in doc.get("comps") or [] if c.get("selected")]
    rows = [["#", "Address", "Dist", "Sold", "Price", "Sqft", "$/sqft", "Bd/Ba", "Built", "Net adj", "Adjusted", "Wt"]]
    for i, c in enumerate(comps, 1):
        a = c.get("adjustment") or {}
        rows.append([str(i), P(_esc(c.get("address")) + ("<br/><font size=6.5 color='#6b6b76'>%s%s</font>" % (_esc(c.get("source")), (", " + _esc(c.get("price_basis"))) if c.get("price_basis") != "closed" else "")), st["p"]),
                     _n(c.get("distance_mi"), nd=2), _esc((c.get("sale_date") or "")[:7]), _m(c.get("sale_price")), _n(c.get("sqft")),
                     _n((c.get("sale_price") or 0) / c["sqft"], nd=0) if c.get("sqft") and c.get("sale_price") else "—",
                     "%s/%s" % (c.get("beds", "—"), c.get("baths", "—")), str(c.get("year_built") or "—"),
                     "%s (%s%%)" % (_m(a.get("net")), a.get("net_pct")), _m(a.get("adjusted_price")), "%s%%" % c.get("weight", "—")])
    t = Table(rows, colWidths=[6 * mm, 48 * mm, 10 * mm, 13 * mm, 17 * mm, 11 * mm, 11 * mm, 11 * mm, 10 * mm, 22 * mm, 17 * mm, 10 * mm], repeatRows=1)
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.2), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f6f5f")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e6e3dc")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (2, 1), (-1, -1), "RIGHT")]))
    story.append(t)
    sel = v.get("selection") or {}
    story.append(P("Searched %.1f mi / %s months; %s candidates, %s qualified, %s used. %s" % (sel.get("radius_used") or 0, sel.get("months_used"), sel.get("candidates"), sel.get("qualified"),
                   rec.get("comps_used"), " ".join(sel.get("notes") or [])), st["small"]))
    rates = v.get("rates") or {}
    if rates:
        story.append(P("Rates: time %s%%/month (%s); living area $%s/sqft (%s); bedroom $%s; full bath $%s; half bath $%s; garage $%s/space; finished basement $%s; pool $%s; lot $%s/sqft; age %s%%/yr; condition %s%% per C-step." % (
            rates.get("market_monthly_pct"), _esc(rates.get("market_basis")), rates.get("gla_per_sqft"), _esc(rates.get("gla_basis")), _n(rates.get("bedroom")), _n(rates.get("bath_full")), _n(rates.get("bath_half")),
            _n(rates.get("garage_space")), _n(rates.get("basement_finished")), _n(rates.get("pool")), rates.get("lot_per_sqft"), rates.get("age_per_year_pct"), rates.get("condition_step_pct")), st["small"]))

    # adjustment grid
    if comps:
        story.append(P("Adjustment grid", st["h2"]))
        keys: List[str] = []
        for c in comps:
            for l in (c.get("adjustment") or {}).get("lines") or []:
                if l["label"] not in keys:
                    keys.append(l["label"])
        grid = [[""] + ["Comp %d" % (i + 1) for i in range(len(comps))]]
        grid.append(["Sale price"] + [_m(c.get("sale_price")) for c in comps])
        for k in keys:
            row = [k]
            for c in comps:
                l = next((x for x in (c.get("adjustment") or {}).get("lines") or [] if x["label"] == k), None)
                row.append(("%+s" % _n(l["amount"])) if l else "")
            grid.append(row)
        grid.append(["Net adjustment"] + ["%s (%s%%)" % (_m((c.get("adjustment") or {}).get("net")), (c.get("adjustment") or {}).get("net_pct")) for c in comps])
        grid.append(["Gross adjustment"] + ["%s (%s%%)" % (_m((c.get("adjustment") or {}).get("gross")), (c.get("adjustment") or {}).get("gross_pct")) for c in comps])
        grid.append(["Adjusted price"] + [_m((c.get("adjustment") or {}).get("adjusted_price")) for c in comps])
        w = min(26, int(150 / max(len(comps), 1)))
        t = Table(grid, colWidths=[32 * mm] + [w * mm] * len(comps), repeatRows=1)
        t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.2), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efede8")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e6e3dc")),
                               ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("BACKGROUND", (0, len(grid) - 1), (-1, len(grid) - 1), colors.HexColor("#e3f1ec")), ("FONTNAME", (0, len(grid) - 1), (-1, len(grid) - 1), "Helvetica-Bold")]))
        story.append(t)
        story.append(P("Reconciliation: weighted mean %s, weighted median %s, spread %s%%; $/sqft indication %s; %s" % (_m(rec.get("weighted_mean")), _m(rec.get("weighted_median")), rec.get("cv_pct"),
                       _m(rec.get("ppsf_indication")), _esc(rec.get("avm_note") or "no AVM blended")), st["small"]))
        story.append(P("Confidence: " + "; ".join(_esc(r) for r in conf.get("reasons") or []), st["small"]))

    # ------------------------------------------------------------ market
    mk = doc.get("market") or {}
    story.append(P("Market", st["h2"]))
    bits = []
    if mk.get("zhvi"):
        z = mk["zhvi"]
        bits.append("Zillow ZHVI for ZIP %s: %s as of %s, %+.1f%% over 12 months, %+.1f%% over 3 months." % (z.get("zip"), _m(z.get("value")), z.get("as_of"), z.get("change_12m_pct") or 0, z.get("change_3m_pct") or 0))
    if mk.get("mortgage"):
        bits.append("30-year fixed %.2f%% (week of %s; a year ago %s%%)." % (mk["mortgage"]["rate_pct"], mk["mortgage"]["as_of"], mk["mortgage"].get("year_ago_pct")))
    if mk.get("flood"):
        bits.append("FEMA flood zone %s — %s" % (mk["flood"].get("zone"), mk["flood"].get("note")))
    if mk.get("acs"):
        a = mk["acs"]
        bits.append("Census tract: median home value %s, median household income %s, owner-occupied %s%%, vacancy %s%%, median year built %s." % (_m(a.get("median_home_value")), _m(a.get("median_household_income")), a.get("owner_share_pct"), a.get("vacancy_pct"), a.get("median_year_built")))
    for g in mk.get("gaps") or []:
        bits.append("Gap: %s." % g)
    story.append(P(" ".join(_esc(b) for b in bits) or "No market context gathered.", st["p"]))
    if an.get("market_read"):
        story.append(P(_md(an["market_read"]), st["p"]))

    # ------------------------------------------------------- rehab & deal
    story.append(P("Rehab budget", st["h2"]))
    story.append(P("%s Total %s (%s/sqft): subtotal %s, location adjustment %s, permits %s, overhead &amp; profit %s, contingency %s." % (_esc(rb.get("basis")), _m(rb.get("total")), _n(rb.get("per_sqft"), nd=1), _m(rb.get("subtotal")),
                   _m(rb.get("location_adjustment")), _m(rb.get("permit_cost")), _m(rb.get("oh_and_p")), _m(rb.get("contingency"))), st["p"]))
    lines = [l for l in rb.get("lines") or [] if l.get("amount")]
    if lines:
        half = (len(lines) + 1) // 2
        cols = [lines[:half], lines[half:]]
        rows = []
        for i in range(half):
            r = []
            for col in cols:
                if i < len(col):
                    r += [col[i]["label"], _m(col[i]["amount"])]
                else:
                    r += ["", ""]
            rows.append(r)
        t = Table(rows, colWidths=[48 * mm, 24 * mm, 48 * mm, 24 * mm])
        t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.5), ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("ALIGN", (3, 0), (3, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, -1), 0.2, colors.HexColor("#e6e3dc"))]))
        story.append(t)
    if an.get("rehab_read"):
        story.append(P(_md(an["rehab_read"]), st["p"]))

    story.append(P("Deal", st["h2"]))
    d = uw.get("deal") or {}
    mao = uw.get("mao") or {}
    a = uw.get("assumptions") or {}
    rows = [["Purchase price (%s)" % _esc(uw.get("price_basis")), _m(d.get("purchase_price")), "MAO, %s%% rule" % _n(mao.get("rule_pct")), _m(mao.get("rule"))],
            ["Rehab (with adders)", _m(uw.get("rehab_total")), "MAO, profit target %s" % _m(mao.get("target_profit")), _m(mao.get("profit_target"))],
            ["Buying costs", _m((uw.get("buying") or {}).get("total")), "Recommended MAO", _m(mao.get("recommended"))],
            ["Holding costs (%s months)" % _n(a.get("months_held")), _m((uw.get("holding") or {}).get("total")), "Profit", _m(d.get("profit"))],
            ["Selling costs", _m((uw.get("selling") or {}).get("total")), "Profit % of ARV", "%s%%" % d.get("profit_pct_of_arv")],
            ["Financing costs", _m((d.get("financing") or {}).get("total")), "Cash needed", _m(d.get("cash_needed"))],
            ["Total project cost", _m(d.get("total_project_cost")), "Cash ROI / annualized", "%s%% / %s%%" % (d.get("cash_roi_pct"), d.get("annualized_roi_pct"))],
            ["Loan (%s)" % _esc((d.get("financing") or {}).get("binding")), _m((d.get("financing") or {}).get("loan")), "LTARV / LTC", "%s%% / %s%%" % ((d.get("financing") or {}).get("ltarv_pct"), (d.get("financing") or {}).get("ltc_pct"))]]
    t = Table(rows, colWidths=[52 * mm, 30 * mm, 52 * mm, 30 * mm])
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("ALIGN", (3, 0), (3, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, -1), 0.2, colors.HexColor("#e6e3dc")),
                           ("FONTNAME", (2, 2), (3, 2), "Helvetica-Bold")]))
    story.append(t)
    for ch in uw.get("checks") or []:
        story.append(P("%s %s: %s (limit %s). %s" % ("✓" if ch.get("ok") else "✗", _esc(ch.get("label")), ch.get("value"), ch.get("limit"), _esc(ch.get("note"))), st["small"]))
    if an.get("mao_reason"):
        story.append(P("<b>Analyst on the offer:</b> %s %s" % (_m(an.get("mao_opinion")), _md(an["mao_reason"])), st["p"]))

    # ----------------------------------------------------------- analyst
    if an:
        story.append(PageBreak())
        story.append(P("Analyst review — %s, effort %s" % (_esc(an.get("model")), _esc(an.get("effort"))), st["h2"]))
        story.append(P("<b>Opinion of value %s (%s – %s), %s confidence.</b> %s" % (_m(an.get("arv_opinion")), _m(an.get("arv_low")), _m(an.get("arv_high")), _esc(an.get("confidence")), _esc(an.get("confidence_reason"))), st["p"]))
        for para in str(an.get("summary") or "").split("\n\n"):
            if para.strip():
                story.append(P(_md(para.strip()), st["p"]))
        if an.get("comp_notes"):
            story.append(P("Comp by comp", st["h2"]))
            idx = {c["comp_id"]: c for c in doc.get("comps") or []}
            for n in an["comp_notes"]:
                c = idx.get(n.get("comp_id")) or {}
                story.append(P("<b>%s</b> — %s: %s" % (_esc(c.get("address") or n.get("comp_id")), _esc(n.get("verdict")), _md(n.get("note") or "")), st["p"]))
        if an.get("adjustment_critique"):
            story.append(P("On the adjustments", st["h2"]))
            for x in an["adjustment_critique"]:
                story.append(P("<b>%s</b> (engine: %s) — %s. %s" % (_esc(x.get("rate")), _esc(x.get("engine_value")), _md(x.get("your_view") or ""), _md(x.get("why") or "")), st["p"]))
        if an.get("risks"):
            story.append(P("Risks", st["h2"]))
            for r in an["risks"]:
                story.append(P("<b>[%s] %s.</b> %s" % (_esc(r.get("severity")), _esc(r.get("title")), _md(r.get("detail") or "")), st["p"]))
        if an.get("questions"):
            story.append(P("Questions that would move the number", st["h2"]))
            for q in an["questions"]:
                story.append(P("<b>%s</b> (%s impact) — %s" % (_esc(q.get("question")), _esc(q.get("impact")), _md(q.get("why_it_matters") or "")), st["p"]))
        if an.get("next_steps"):
            story.append(P("Next steps", st["h2"]))
            for s in an["next_steps"]:
                story.append(P("• " + _md(s), st["p"]))
    story.append(Spacer(1, 10))
    story.append(P("Sources: %s. Valuation is an automated sales-comparison opinion for underwriting support, not an appraisal under USPAP. Comps marked 'web' were found on public listing pages and are unverified against MLS or county records." % _esc(", ".join(doc.get("data_sources") or [])), st["small"]))
    pdoc.build(story)
    return buf.getvalue()


# ================================================================= XLSX
def xlsx(doc: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "MAO Rehab Calculator"
    uw = doc.get("underwriting") or {}
    a = uw.get("assumptions") or {}
    v = doc.get("valuation") or {}
    rb = doc.get("rehab_budget") or {}
    geo = doc.get("geocode") or {}
    inp = PatternFill("solid", fgColor="FFF8DC")
    hdr = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F6F5F")
    blue = Font(color="1F4ED8")
    money = '"$"#,##0'
    pct = "0.0%"

    def H(r, c, text, span=2):
        ws.cell(r, c, text).font = hdr
        ws.cell(r, c).fill = hfill
        for k in range(1, span):
            ws.cell(r, c + k).fill = hfill

    def I(r, c, label, value, fmt=money):   # input cell
        ws.cell(r, c, label)
        cell = ws.cell(r, c + 1, value)
        cell.fill = inp
        cell.font = blue
        cell.number_format = fmt
        return cell.coordinate

    def F(r, c, label, formula, fmt=money, bold=False):   # formula cell
        ws.cell(r, c, label).font = Font(bold=bold)
        cell = ws.cell(r, c + 1, formula)
        cell.number_format = fmt
        cell.font = Font(bold=bold)
        return cell.coordinate

    ws.cell(1, 1, "MAO Rehab Calculator — " + str(geo.get("formatted") or (doc.get("subject") or {}).get("address") or "")).font = Font(bold=True, size=13)
    ws.cell(2, 1, "Blue on yellow = input · black = formula · generated by MangoTree ARV %s" % datetime.now(timezone.utc).strftime("%Y-%m-%d")).font = Font(italic=True, color="6B6B76")
    H(4, 1, "STEP 1: AFTER REPAIR VALUE")
    arv = I(5, 1, "ARV", v.get("arv") or 0)
    H(7, 1, "STEP 2: FIXED COSTS")
    ws.cell(8, 1, "Buying costs").font = Font(bold=True)
    b = uw.get("buying") or {}
    r = 9
    buy_cells = []
    for lab, key in (("Brokerage fee", "brokerage"), ("Inspection", "inspection"), ("Appraisal", "appraisal"), ("Title search", "title_search"), ("Title insurance", "title_insurance"), ("Other", "other")):
        buy_cells.append(I(r, 1, lab, b.get(key) or 0)); r += 1
    buy_total = F(r, 1, "Total buying costs", "=SUM(%s:%s)" % (buy_cells[0], buy_cells[-1]), bold=True); r += 2
    ws.cell(r, 1, "Holding costs").font = Font(bold=True); r += 1
    months = I(r, 1, "Months held", a.get("months_held") or 8, "0.0"); r += 1
    h = uw.get("holding") or {}
    taxes = I(r, 1, "Property taxes (annual)", h.get("taxes_annual_used") or 0); r += 1
    ins = I(r, 1, "Insurance (annual)", a.get("insurance_annual") or 0); r += 1
    util = I(r, 1, "Utilities (monthly)", a.get("utilities_monthly") or 0); r += 1
    hoa = I(r, 1, "HOA (monthly)", a.get("hoa_monthly") or 0); r += 1
    maint = I(r, 1, "Maintenance (monthly)", a.get("maintenance_monthly") or 0); r += 1
    hother = I(r, 1, "Other (monthly)", a.get("holding_other_monthly") or 0); r += 1
    hold_total = F(r, 1, "Total holding costs", "=(%s/12+%s/12+%s+%s+%s+%s)*%s" % (taxes, ins, util, hoa, maint, hother, months), bold=True); r += 2
    ws.cell(r, 1, "Selling costs").font = Font(bold=True); r += 1
    s = uw.get("selling") or {}
    comm = I(r, 1, "Commission %", (a.get("commission_pct") or 0) / 100, pct); r += 1
    assist = I(r, 1, "Seller-assisted closing %", (a.get("seller_assist_pct") or 0) / 100, pct); r += 1
    ttax = I(r, 1, "Transfer tax %", (s.get("transfer_tax_pct") or 0) / 100, pct); r += 1
    sb = I(r, 1, "Brokerage fee", a.get("sell_brokerage") or 0); r += 1
    hw = I(r, 1, "Home warranty", a.get("home_warranty") or 0); r += 1
    sti = I(r, 1, "Title insurance", a.get("sell_title_insurance") or 0); r += 1
    app2 = I(r, 1, "2nd appraisal", a.get("second_appraisal") or 0); r += 1
    sother = I(r, 1, "Other", a.get("sell_other") or 0); r += 1
    sell_total = F(r, 1, "Total selling costs", "=%s*(%s+%s+%s)+%s+%s+%s+%s+%s" % (arv, comm, assist, ttax, sb, hw, sti, app2, sother), bold=True); r += 2
    ws.cell(r, 1, "Financing").font = Font(bold=True); r += 1
    rate = I(r, 1, "Annual interest rate", (a.get("rate_pct") or 0) / 100, pct); r += 1
    points = I(r, 1, "Points", (a.get("points_pct") or 0) / 100, pct); r += 1
    fees = I(r, 1, "Loan fees", a.get("loan_fees") or 0); r += 1
    ltc = I(r, 1, "Max loan-to-cost", (a.get("ltc_pct") or 0) / 100, pct); r += 1
    ltarv = I(r, 1, "Max loan-to-ARV", (a.get("ltarv_pct") or 0) / 100, pct); r += 1
    fin_row = r
    r += 3

    # repairs block (columns D:E)
    H(4, 4, "STEP 3: REPAIR COSTS")
    rr = 5
    line_cells = []
    for l in rb.get("lines") or []:
        line_cells.append(I(rr, 4, l["label"], l.get("amount") or 0)); rr += 1
    sub = F(rr, 4, "Subtotal estimate", "=SUM(%s:%s)" % (line_cells[0], line_cells[-1]) if line_cells else "=0", bold=True); rr += 1
    ad = rb.get("adders") or {}
    locm = I(rr, 4, "Location multiplier", ad.get("location_multiplier") or 1, "0.00"); rr += 1
    permit = I(rr, 4, "Permit cost %", (ad.get("permit_pct") or 0) / 100, pct); rr += 1
    ohp = I(rr, 4, "Contractor OH & P %", (ad.get("oh_and_p_pct") or 0) / 100, pct); rr += 1
    cont = I(rr, 4, "Contingency %", (ad.get("contingency_pct") or 0) / 100, pct); rr += 1
    adders = F(rr, 4, "Total adders", "=%s*(%s-1)+%s*%s*(%s+%s+%s)" % (sub, locm, sub, locm, permit, ohp, cont)); rr += 1
    repairs = F(rr, 4, "TOTAL REPAIR ESTIMATE", "=%s+%s" % (sub, adders), bold=True); rr += 2

    # purchase / results block (columns G:H)
    H(4, 7, "PURCHASE PRICE")
    price = I(5, 7, "Purchase price", (uw.get("deal") or {}).get("purchase_price") or 0)
    H(7, 7, "70% RULE")
    rule = I(8, 7, "Rule %", (a.get("rule_pct") or 70) / 100, pct)
    mao = F(9, 7, "Maximum allowable offer", "=%s*%s-%s" % (arv, rule, repairs), bold=True)
    H(11, 7, "FINANCING")
    loan = F(12, 7, "Loan amount", "=MIN((%s+%s)*%s,%s*%s)" % (price, repairs, ltc, arv, ltarv))
    fin_total = F(13, 7, "Total financing costs", "=%s*%s*%s/12+%s*%s+%s" % (loan, rate, months, loan, points, fees), bold=True)
    ws.cell(fin_row, 1, "Total financing costs (see H13)")
    H(15, 7, "CALCULATED REHAB PROFIT")
    fixed = F(16, 7, "Total fixed costs", "=%s+%s+%s+%s" % (buy_total, hold_total, sell_total, fin_total))
    total = F(17, 7, "Total project costs", "=%s+%s+%s" % (price, repairs, fixed), bold=True)
    profit = F(18, 7, "Profit", "=%s-%s" % (arv, total), bold=True)
    F(19, 7, "Profit per month", "=%s/%s" % (profit, months))
    F(20, 7, "Profit % of ARV", "=%s/%s" % (profit, arv), pct)
    H(22, 7, "INVESTMENT BREAKDOWN")
    F(23, 7, "Capital needed", "=%s" % total)
    F(24, 7, "Debt", "=%s" % loan)
    cash = F(25, 7, "Cash (equity)", "=%s-%s" % (total, loan))
    F(26, 7, "Cash ROI", "=IF(%s>0,%s/%s,0)" % (cash, profit, cash), pct)
    F(27, 7, "Annualized ROI", "=IF(%s>0,%s/%s*12/%s,0)" % (cash, profit, cash, months), pct)
    for col, w in ((1, 30), (2, 14), (4, 28), (5, 14), (7, 28), (8, 14)):
        ws.column_dimensions[get_column_letter(col)].width = w

    # comps sheet
    wc = wb.create_sheet("Comps")
    head = ["#", "Used", "Source", "Address", "Distance (mi)", "Sale date", "Sale price", "Price basis", "Sqft", "$/sqft", "Beds", "Baths", "Year built", "Lot sqft", "Condition", "Similarity", "Net adj", "Net %", "Gross %", "Adjusted price", "Weight %", "URL"]
    wc.append(head)
    for c in wc[1]:
        c.font = hdr
        c.fill = hfill
    for i, c in enumerate(doc.get("comps") or [], 1):
        adj = c.get("adjustment") or {}
        wc.append([i, "yes" if c.get("selected") else "", c.get("source"), c.get("address"), c.get("distance_mi"), c.get("sale_date"), c.get("sale_price"), c.get("price_basis"), c.get("sqft"),
                   round(c["sale_price"] / c["sqft"], 1) if c.get("sqft") and c.get("sale_price") else None, c.get("beds"), c.get("baths"), c.get("year_built"), c.get("lot_sqft"), c.get("condition"),
                   c.get("similarity"), adj.get("net"), adj.get("net_pct"), adj.get("gross_pct"), adj.get("adjusted_price"), c.get("weight"), c.get("url")])
    for col in wc.columns:
        wc.column_dimensions[col[0].column_letter].width = max(10, min(50, max(len(str(x.value or "")) for x in col) + 2))
    wc.freeze_panes = "A2"

    # adjustments sheet
    wa = wb.create_sheet("Adjustments")
    wa.append(["Comp", "Address", "Line", "Amount", "Why"])
    for c in wa[1]:
        c.font = hdr
        c.fill = hfill
    for i, c in enumerate([c for c in doc.get("comps") or [] if c.get("selected")], 1):
        for l in (c.get("adjustment") or {}).get("lines") or []:
            wa.append([i, c.get("address"), l["label"], l["amount"], l["why"]])
    wa.column_dimensions["B"].width = 40
    wa.column_dimensions["C"].width = 24
    wa.column_dimensions["E"].width = 80
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
